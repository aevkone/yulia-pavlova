# Сценарии цветочной рубрики → .docx: каждый сценарий – цветной блок, метки выделены цветом
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

SRC = 'C:/Users/buhsk/UlyaPavlova/docs/Сценарии_цветочный_бизнес.md'
OUT = 'C:/Users/buhsk/UlyaPavlova/docs/Сценарии_цветочный_бизнес.docx'
HEAD, BODY = 'Oswald', 'Manrope'
# (акцент, фон блока)
PALETTE = [('C0392B', 'FDECEA'), ('C9A24B', 'FBF5E6'), ('2E86AB', 'E8F3F8'), ('8E44AD', 'F4ECF8'),
           ('27AE60', 'E9F7EF'), ('D35400', 'FDF0E6')]
LABELS = {'Хук': '🎬', 'Ситуация': '📍', 'Версии': '🤔', 'Версии / три урока:': '🤔', 'Вопрос': '💬',
          'Вывод': '✅', 'Мост': '👉', 'Подпись к ролику (начало):': '✍️'}


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_)
    tcPr.append(sh)


def borders(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side, sz in (('left', 36), ('top', 4), ('bottom', 4), ('right', 4)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), hex_)
        b.append(e)
    tcPr.append(b)


def run(p, text, font=BODY, size=10.5, bold=False, italic=False, color=None):
    r = p.add_run(text); r.font.name = font; r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r


def inline(p, text, color, size=10.5):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if part:
            run(p, part.replace('*', ''), bold=bool(i % 2), color=color if i % 2 else None, size=size)


def para(container, space_after=4):
    p = container.add_paragraph(); p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    return p


md = open(SRC, encoding='utf-8').read()
parts = re.split(r'\n(?=## )', md)
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)

# Шапка
intro = parts[0].splitlines()
p = para(doc, 2); run(p, 'ИСТОРИИ ЦВЕТОЧНОГО БИЗНЕСА', HEAD, 26, True, color='C0392B')
p = para(doc, 10); run(p, '18 сценариев роликов · Юлия Павлова · рубрика для споров в комментариях', BODY, 11, color='777777')
t = doc.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, 'F6F6F6'); borders(c, '444444')
c.paragraphs[0].text = ''
first = True
for l in intro[1:]:
    l = l.strip()
    if not l or l == '---':
        continue
    l = l.lstrip('> ').strip()
    p = c.paragraphs[0] if first else c.add_paragraph(); first = False
    p.paragraph_format.space_after = Pt(4)
    inline(p, l, '444444', 10)
para(doc, 6)

n = 0
for sec in parts[1:]:
    lines = sec.splitlines()
    title = lines[0][3:].strip()
    body = [l for l in lines[1:]]
    m = re.match(r'(\d+)\. (.+)', title)
    if not m:  # служебные разделы: «Ещё 10 идей», «Порядок публикации»
        p = para(doc, 6); p.paragraph_format.space_before = Pt(10)
        run(p, title.upper(), HEAD, 16, True, color='222222')
        rows = [l for l in body if l.startswith('|') and not re.match(r'^\|[-| ]+\|$', l)]
        if rows:
            cells = [[x.strip() for x in r.strip('|').split('|')] for r in rows]
            tb = doc.add_table(rows=len(cells), cols=3); tb.style = 'Table Grid'
            for i, row in enumerate(cells):
                for j, v in enumerate(row):
                    cc = tb.cell(i, j); cc.paragraphs[0].text = ''
                    inline(cc.paragraphs[0], v, None, 9.5)
                    if i == 0:
                        shade(cc, '222222'); cc.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255); cc.paragraphs[0].runs[0].bold = True
                    elif i % 2 == 0:
                        shade(cc, 'F6F6F6')
            for row in tb.rows:
                row.cells[0].width = Cm(1.4); row.cells[1].width = Cm(6); row.cells[2].width = Cm(10)
        for l in body:
            l = l.strip()
            if l and not l.startswith('|') and l != '---':
                p = para(doc, 4); inline(p, l.lstrip('> ').strip(), '444444', 10)
                if l.startswith('>'):
                    for r in p.runs: r.italic = True
        continue
    acc, bg = PALETTE[n % len(PALETTE)]; n += 1
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.cell(0, 0); shade(c, bg); borders(c, acc)
    trPr = tb.rows[0]._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(6)
    run(p, f'{m.group(1)}  ', HEAD, 22, True, color=acc)
    run(p, m.group(2), HEAD, 15, True, color='222222')
    for l in body:
        l = l.rstrip()
        if not l.strip() or l.strip() == '---':
            continue
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        lab = re.match(r'^\*\*(.+?)\*\*\s*(.*)$', l.strip())
        base = re.split(r'[ (.:/]', lab.group(1))[0] if lab else ''
        if lab and base in ('Хук', 'Ситуация', 'Версии', 'Вопрос', 'Вывод', 'Мост', 'Подпись'):
            key = lab.group(1).rstrip('.:')
            em = {'Хук': '🎬', 'Ситуация': '📍', 'Версии': '🤔', 'Вопрос': '💬', 'Вывод': '✅', 'Мост': '👉', 'Подпись': '✍️'}[base]
            run(p, f'{em} {key.upper()}  ', HEAD, 10.5, True, color=acc)
            inline(p, lab.group(2), acc)
        elif l.lstrip().startswith('–') or re.match(r'^\d\. ', l.strip()):
            p.paragraph_format.left_indent = Cm(0.6)
            inline(p, l.strip(), acc)
        elif l.startswith('>'):
            inline(p, l.lstrip('> ').strip(), acc, 9.5)
            for r in p.runs: r.italic = True; r.font.color.rgb = RGBColor.from_string('666666')
        else:
            inline(p, l.strip(), acc)
    para(doc, 8)

doc.save(OUT)
print('ok', n)
